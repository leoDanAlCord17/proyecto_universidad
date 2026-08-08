# -*- coding: utf-8 -*-
"""
Genera Plan_de_Pruebas_Manuales_QA.docx — plan de pruebas manuales de QA
para Activiti (Flutter Web/PWA + Supabase), cubriendo el 100% de las
pantallas/rutas de la app.

Uso:
    python generate_docs.py

Requiere: python-docx (pip install python-docx)
"""

from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ─── Paleta corporativa (tomada de lib/configuracion/colores_app.dart) ──────

COLOR_ACENTO = RGBColor(0x5B, 0x3F, 0xD4)
COLOR_ACENTO_CLARO = "EDE9FB"  # hex sin '#' para shading de celdas
COLOR_TEXTO = RGBColor(0x22, 0x1F, 0x2B)
COLOR_GRIS = RGBColor(0x6B, 0x6B, 0x76)

FECHA_HOY = date.today().strftime("%d/%m/%Y")


# ─── Helpers de estilo (python-docx no cubre shading/bordes de forma nativa) ─

def sombrear_celda(celda, color_hex):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def bordes_celda(celda, color_hex="D9D5E8", size=4):
    tcPr = celda._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color_hex)
        borders.append(el)
    tcPr.append(borders)


def texto_celda(celda, texto, negrita=False, color=None, tamanio=10):
    celda.text = ""
    parrafo = celda.paragraphs[0]
    for i, linea in enumerate(str(texto).split("\n")):
        if i > 0:
            parrafo.add_run().add_break()
        run = parrafo.add_run(linea)
        run.bold = negrita
        run.font.size = Pt(tamanio)
        if color:
            run.font.color.rgb = color


def agregar_encabezado_portada(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ACTIVITI")
    run.font.size = Pt(16)
    run.bold = True
    run.font.color.rgb = COLOR_ACENTO

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("Plan de Pruebas Manuales — QA")
    run.font.size = Pt(28)
    run.bold = True
    run.font.color.rgb = COLOR_TEXTO
    doc.add_paragraph()

    meta = [
        ("Versión de la prueba:", "1.0"),
        ("Fecha:", FECHA_HOY),
        ("Alcance:", "Revisión funcional y visual completa (100% de pantallas y flujos)"),
        ("Plataforma:", "Flutter Web / PWA (móvil y navegador)"),
        ("Backend:", "Supabase (Auth, Postgres/PostgREST, Realtime, Edge Functions, FCM)"),
    ]
    tabla = doc.add_table(rows=len(meta), cols=2)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    for i, (etiqueta, valor) in enumerate(meta):
        fila = tabla.rows[i]
        fila.cells[0].width = Cm(4.5)
        fila.cells[1].width = Cm(10)
        texto_celda(fila.cells[0], etiqueta, negrita=True, color=COLOR_ACENTO)
        texto_celda(fila.cells[1], valor)
        for c in fila.cells:
            bordes_celda(c)
    doc.add_page_break()


def agregar_instrucciones(doc):
    h = doc.add_heading("Instrucciones para el QA Tester", level=1)
    for run in h.runs:
        run.font.color.rgb = COLOR_ACENTO

    intro = doc.add_paragraph(
        "Antes de comenzar, prepara el entorno de prueba siguiendo esta lista. "
        "Marca cada punto una vez confirmado."
    )
    intro.runs[0].font.color.rgb = COLOR_GRIS

    puntos = [
        "Ten conexión estable a internet (Wi-Fi y datos móviles) — Activiti depende de Supabase en tiempo real, "
        "no funciona completamente offline salvo caché ya cargado.",
        "Prueba en al menos dos tamaños de pantalla: un teléfono real (ideal, para permisos de cámara/push reales) "
        "y un navegador de escritorio con DevTools en modo responsive.",
        "Ten a mano credenciales de al menos dos roles distintos: un usuario 'estudiante' sin permisos "
        "administrativos y un usuario 'administrador' con permisos completos — varios flujos solo son visibles "
        "para uno u otro.",
        "Ten un evento de prueba disponible (no uses eventos reales con estudiantes) para los casos de "
        "creación/edición/escaneo/cierre, para no afectar datos reales.",
        "Verifica los permisos del dispositivo antes de empezar: cámara (para escaneo QR) y notificaciones "
        "push — prueba también qué pasa cuando NIEGAS esos permisos a propósito (ver casos de borde).",
        "Prueba la orientación de pantalla (vertical y horizontal) al menos en las pantallas con formularios "
        "largos y en el escáner de QR.",
        "Si la app está instalada como PWA (ícono en el escritorio/pantalla de inicio), prueba también el gesto "
        "de 'atrás' del sistema y el deslizar desde el borde — no debe salir de la app ni mostrar una pantalla en blanco.",
        "Anota TODO lo que veas raro en la columna de Observaciones, aunque no estés seguro de si es un bug — "
        "es más fácil descartarlo después que perderlo.",
        "Al terminar cada módulo, pasa los hallazgos importantes a la sección final "
        "'Resumen de Hallazgos y Errores Encontrados'.",
    ]
    for punto in puntos:
        p = doc.add_paragraph(punto, style="List Bullet")
        p.paragraph_format.space_after = Pt(6)

    doc.add_paragraph()
    leyenda = doc.add_paragraph()
    run = leyenda.add_run(
        "Cómo leer cada caso de prueba: cada caso se presenta como una ficha con 8 campos — "
        "ID, Módulo/Pantalla, Nombre/Propósito, Precondiciones, Pasos, Resultado Esperado, "
        "Estado de Verificación (marca ☐ Pasó o ☐ Falló) y Observaciones (para anotar a mano)."
    )
    run.italic = True
    run.font.color.rgb = COLOR_GRIS
    doc.add_page_break()


def agregar_indice_modulos(doc, modulos):
    h = doc.add_heading("Índice de módulos cubiertos", level=1)
    for run in h.runs:
        run.font.color.rgb = COLOR_ACENTO
    for nombre, casos in modulos:
        total = len(casos)
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(f"{nombre} ")
        run.bold = True
        p.add_run(f"— {total} caso{'s' if total != 1 else ''} de prueba")
    doc.add_page_break()


def agregar_caso_de_prueba(doc, caso):
    campos = [
        ("ID Caso de Prueba", caso["id"]),
        ("Módulo / Pantalla", caso["pantalla"]),
        ("Nombre / Propósito", caso["nombre"]),
        ("Precondiciones", caso["precondiciones"]),
        (
            "Pasos a ejecutar",
            "\n".join(f"{i + 1}. {paso}" for i, paso in enumerate(caso["pasos"])),
        ),
        ("Resultado Esperado", caso["esperado"]),
        ("Estado de Verificación", "☐  Pasó          ☐  Falló"),
        ("Observaciones", " "),
    ]

    tabla = doc.add_table(rows=len(campos), cols=2)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    for i, (etiqueta, valor) in enumerate(campos):
        fila = tabla.rows[i]
        celda_etq, celda_val = fila.cells
        celda_etq.width = Cm(4.2)
        celda_val.width = Cm(11.8)

        es_id = etiqueta == "ID Caso de Prueba"
        texto_celda(
            celda_etq,
            etiqueta,
            negrita=True,
            color=RGBColor(0xFF, 0xFF, 0xFF) if es_id else COLOR_ACENTO,
            tamanio=10,
        )
        sombrear_celda(celda_etq, "5B3FD4" if es_id else COLOR_ACENTO_CLARO)
        bordes_celda(celda_etq)

        negrita_val = es_id or etiqueta == "Resultado Esperado"
        texto_celda(celda_val, valor, negrita=negrita_val, tamanio=10)
        if etiqueta == "Observaciones":
            # Deja espacio en blanco físico para escribir a mano.
            for _ in range(3):
                celda_val.add_paragraph()
        bordes_celda(celda_val)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def agregar_modulo(doc, nombre, casos):
    h = doc.add_heading(nombre, level=1)
    for run in h.runs:
        run.font.color.rgb = COLOR_ACENTO
    for caso in casos:
        agregar_caso_de_prueba(doc, caso)
    doc.add_page_break()


def agregar_resumen_hallazgos(doc):
    h = doc.add_heading("Resumen de Hallazgos y Errores Encontrados", level=1)
    for run in h.runs:
        run.font.color.rgb = COLOR_ACENTO

    p = doc.add_paragraph(
        "Usa esta tabla para consolidar los bugs, comportamientos anómalos o detalles visuales "
        "detectados durante la ejecución del plan completo. Referencia el ID del caso de prueba "
        "correspondiente cuando aplique."
    )
    p.runs[0].font.color.rgb = COLOR_GRIS

    encabezados = ["ID Caso", "Pantalla", "Descripción del hallazgo", "Severidad\n(Crítica/Alta/Media/Baja)", "Estado\n(Abierto/Resuelto)"]
    filas_vacias = 15
    tabla = doc.add_table(rows=1 + filas_vacias, cols=len(encabezados))
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    encabezado = tabla.rows[0]
    for i, texto in enumerate(encabezados):
        celda = encabezado.cells[i]
        texto_celda(celda, texto, negrita=True, color=RGBColor(0xFF, 0xFF, 0xFF), tamanio=9)
        sombrear_celda(celda, "5B3FD4")
        bordes_celda(celda)

    for f in range(1, filas_vacias + 1):
        for c in range(len(encabezados)):
            celda = tabla.rows[f].cells[c]
            celda.text = ""
            if f % 2 == 0:
                sombrear_celda(celda, "F5F3FB")
            bordes_celda(celda)

    doc.add_paragraph()
    firma = doc.add_paragraph()
    firma.add_run("Ejecutado por: ").bold = True
    firma.add_run("_______________________________          ")
    firma.add_run("Fecha: ").bold = True
    firma.add_run("______________")


# ─── Datos: casos de prueba por módulo ──────────────────────────────────────
# Basado en el árbol de rutas real (lib/compartido/constantes.dart → class Rutas,
# lib/configuracion/rutas/*.dart) y las pantallas correspondientes en
# lib/funcionalidades/**/*_pantalla.dart — cobertura de las 34 rutas de la app
# (excluyendo /dev/widgets y /dev/fuentes, marcadas en el código como
# exclusivas de desarrollo, a eliminar antes de producción).

MODULOS = []


def modulo(nombre, casos):
    MODULOS.append((nombre, casos))


def caso(id_, pantalla, nombre, precondiciones, pasos, esperado):
    return {
        "id": id_,
        "pantalla": pantalla,
        "nombre": nombre,
        "precondiciones": precondiciones,
        "pasos": pasos,
        "esperado": esperado,
    }


# 1. AUTENTICACIÓN Y PERFIL ---------------------------------------------------
modulo("1. Autenticación y Perfil", [
    caso(
        "TC-AUTH-001", "Splash (/)", "Redirección automática según sesión",
        "Ninguna — app recién abierta",
        [
            "Abre la app (URL raíz) sin sesión activa.",
            "Observa la pantalla de carga inicial (splash).",
            "Espera a que redirija automáticamente.",
        ],
        "Si no hay sesión, redirige a Login. Si hay sesión válida, redirige directo a Inicio sin pasar por Login. "
        "El splash no debe quedarse colgado más de unos segundos.",
    ),
    caso(
        "TC-AUTH-002", "Login (/login)", "Inicio de sesión exitoso (happy path)",
        "Cuenta aprobada y activa existente",
        [
            "Ingresa un correo válido y su contraseña correcta.",
            "Presiona 'Iniciar sesión'.",
        ],
        "Muestra un indicador de carga breve, luego navega a Inicio (Home) mostrando el saludo con el nombre "
        "del usuario correcto. No debe quedar residuo del formulario de login en el historial de navegación "
        "(el botón atrás no debe volver al login).",
    ),
    caso(
        "TC-AUTH-003", "Login (/login)", "Credenciales incorrectas",
        "Ninguna",
        [
            "Ingresa un correo válido pero una contraseña incorrecta.",
            "Presiona 'Iniciar sesión'.",
        ],
        "Muestra el mensaje 'Correo o contraseña incorrectos.' sin navegar de pantalla, y el botón vuelve a "
        "estar disponible (no se queda cargando indefinidamente).",
    ),
    caso(
        "TC-AUTH-004", "Login (/login)", "Campos vacíos",
        "Ninguna",
        [
            "Deja ambos campos vacíos.",
            "Presiona 'Iniciar sesión'.",
        ],
        "Muestra 'Por favor, llena todos los campos.' sin llegar a llamar al servidor (no debe verse un spinner "
        "de carga de red).",
    ),
    caso(
        "TC-AUTH-005", "Login (/login)", "Formato de correo inválido",
        "Ninguna",
        [
            "Escribe un texto sin formato de correo (ej. 'abc123') en el campo de correo.",
            "Escribe cualquier contraseña.",
            "Presiona 'Iniciar sesión'.",
        ],
        "Muestra 'Ingresa un correo con formato válido.' antes de intentar conectarse al servidor.",
    ),
    caso(
        "TC-AUTH-006", "Login (/login)", "Reintento automático ante error transitorio del servidor",
        "Requiere simular inestabilidad de red (modo avión intermitente o throttling en DevTools)",
        [
            "Activa un throttling de red muy lento o desconecta brevemente la red justo al presionar 'Iniciar sesión'.",
            "Reconecta la red a los 1-2 segundos.",
            "Observa si el login se completa solo, sin que tengas que tocar el botón de nuevo.",
        ],
        "La app reintenta automáticamente una vez ante un fallo transitorio del servidor de autenticación; "
        "si la red vuelve a tiempo, el login se completa sin intervención manual.",
    ),
    caso(
        "TC-AUTH-007", "Registro (/registro)", "Registro de cuenta nueva",
        "Correo no usado previamente en el sistema",
        [
            "Desde Login, presiona el enlace para crear cuenta.",
            "Completa los 3 campos requeridos con datos válidos.",
            "Presiona 'Continuar'.",
        ],
        "Crea el usuario en Auth y navega al flujo de completar perfil. No debe permitir continuar con campos vacíos.",
    ),
    caso(
        "TC-AUTH-008", "Registro (/registro)", "Correo ya registrado",
        "Correo que ya tiene cuenta creada",
        [
            "Intenta registrarte con un correo que ya existe en el sistema.",
        ],
        "Muestra 'Este correo ya tiene una cuenta. Inicia sesión.' y no crea una cuenta duplicada.",
    ),
    caso(
        "TC-AUTH-009", "Completar perfil (/completar_perfil)", "Completar perfil tras registro",
        "Cuenta creada en Auth pero sin perfil en la tabla de usuarios",
        [
            "Completa nombre(s), apellido(s) y demás campos solicitados.",
            "Envía el formulario.",
        ],
        "Guarda el perfil y navega a 'Pendiente de aprobación' (si la revisión de cuentas está habilitada) o "
        "directo a Inicio (si no lo está).",
    ),
    caso(
        "TC-AUTH-010", "Pendiente aprobación (/pendiente_aprobacion)", "Cuenta esperando revisión",
        "Cuenta con perfil completo pero estatus_aprobacion = pendiente",
        [
            "Inicia sesión con una cuenta recién registrada, aún no aprobada por un administrador.",
        ],
        "Muestra una pantalla clara indicando que la cuenta está en revisión, sin acceso al resto de la app. "
        "No debe permitir navegar manualmente (cambiando la URL) a rutas internas.",
    ),
    caso(
        "TC-AUTH-011", "Usuario rechazado (/usuario_rechazado)", "Cuenta rechazada por un administrador",
        "Cuenta con estatus_aprobacion = rechazado",
        [
            "Inicia sesión con una cuenta que un administrador rechazó previamente.",
        ],
        "Muestra la pantalla de cuenta rechazada con un mensaje claro, sin acceso al resto de la app.",
    ),
    caso(
        "TC-AUTH-012", "Recuperar contraseña (/recuperar_contrasena)", "Solicitar enlace de recuperación",
        "Ninguna",
        [
            "Desde Login, presiona '¿Olvidaste tu contraseña?'.",
            "Ingresa un correo (exista o no en el sistema).",
            "Envía la solicitud.",
        ],
        "Siempre muestra un mensaje de éxito genérico (por seguridad no revela si el correo existe o no). "
        "Si el correo existe, debe llegar un email con el enlace de recuperación.",
    ),
    caso(
        "TC-AUTH-013", "Nueva contraseña (/nueva_contrasena)", "Establecer nueva contraseña desde el enlace del correo",
        "Enlace de recuperación válido y reciente",
        [
            "Abre el enlace de recuperación recibido por correo.",
            "Ingresa una nueva contraseña que cumpla los requisitos.",
            "Confirma el cambio.",
        ],
        "Actualiza la contraseña y permite iniciar sesión con la nueva clave. El enlace no debe poder reutilizarse "
        "una segunda vez.",
    ),
    caso(
        "TC-AUTH-014", "Sesión / cualquier pantalla", "Sesión iniciada en otro dispositivo",
        "Misma cuenta con sesión activa en dos dispositivos/navegadores distintos",
        [
            "Inicia sesión con la misma cuenta en el Dispositivo A.",
            "Inicia sesión con la misma cuenta en el Dispositivo B.",
            "Vuelve al Dispositivo A y realiza cualquier acción.",
        ],
        "El Dispositivo A muestra un aviso 'Tu sesión fue iniciada en otro dispositivo' (aviso flotante visible, "
        "sin necesidad de refrescar la página) y cierra la sesión.",
    ),
    caso(
        "TC-AUTH-015", "Sesión / cualquier pantalla", "Expiración natural de la sesión (no por otro dispositivo)",
        "Sesión iniciada, sin usar la app activamente por un período largo (más de 1 hora) o con la hora del "
        "dispositivo adelantada para forzar el vencimiento del token",
        [
            "Inicia sesión y deja la app abierta e inactiva (sin tocarla) durante más de una hora, o adelanta "
            "manualmente la hora del dispositivo.",
            "Vuelve a interactuar con la app (navega a otra pantalla o realiza una acción que llame al servidor).",
        ],
        "La app detecta el token vencido y, o bien lo renueva automáticamente de forma transparente sin que el "
        "usuario note nada, o lo redirige a Login con un mensaje claro de sesión expirada — nunca debe quedar "
        "en un estado roto (pantalla en blanco, errores de red repetidos, o datos a medio cargar sin explicación).",
    ),
])

# 2. NAVEGACIÓN Y DASHBOARD ---------------------------------------------------
modulo("2. Navegación y Dashboard", [
    caso(
        "TC-NAV-001", "Navegación principal", "Cambio entre pestañas del menú inferior",
        "Sesión iniciada",
        [
            "Toca cada ícono de la barra inferior: Inicio, Eventos, Escanear, Historial, Perfil.",
            "Regresa a una pestaña ya visitada.",
        ],
        "Cada pestaña carga su contenido correctamente. Al volver a una pestaña ya visitada, conserva su estado "
        "(no vuelve a mostrar el spinner de carga inicial ni pierde el scroll).",
    ),
    caso(
        "TC-NAV-002", "Navegación principal", "Botón 'Escanear' abre pantalla completa, no una pestaña",
        "Sesión iniciada",
        [
            "Toca el ícono 'Escanear' en la barra inferior.",
            "Presiona el botón atrás del sistema (o gesto de deslizar) para salir.",
        ],
        "Escanear se abre como pantalla superpuesta (no queda seleccionada como pestaña activa). Al salir, "
        "regresa a la pestaña donde estabas antes, no al Home por defecto.",
    ),
    caso(
        "TC-NAV-003", "Navegación principal (PWA instalada)", "Gesto de 'atrás' del sistema no rompe la app",
        "App instalada como PWA en Android",
        [
            "Con la app instalada, navega entre 2-3 pestañas del menú inferior.",
            "Desliza desde el borde izquierdo de la pantalla (gesto de 'atrás' de Android).",
            "Presiona el botón físico/gesto de 'atrás' del sistema varias veces seguidas.",
        ],
        "El gesto no saca de la app ni muestra una pantalla en blanco/anterior inesperada. La app debe "
        "permanecer en un estado coherente.",
    ),
    caso(
        "TC-NAV-004", "Navegación / rutas protegidas", "Restricción de rutas por permiso",
        "Cuenta sin permisos administrativos",
        [
            "Con una cuenta de estudiante (sin permisos de administración), intenta escribir manualmente en la "
            "barra de direcciones una ruta administrativa (ej. /gestion_usuarios o /permisos).",
        ],
        "Redirige automáticamente a Inicio en vez de mostrar la pantalla protegida.",
    ),
    caso(
        "TC-NAV-005", "Navegación / recarga", "Recarga de página (F5) conserva la sesión y la ruta",
        "Sesión iniciada, en cualquier pantalla interna",
        [
            "Navega a una pantalla interna (ej. Eventos).",
            "Recarga el navegador (F5) o refresca el PWA.",
        ],
        "Vuelve a cargar sin cerrar sesión y, si es razonable, regresa a la misma ubicación en vez de reiniciar "
        "siempre en Inicio.",
    ),
])

# 3. INICIO (HOME) -------------------------------------------------------------
modulo("3. Inicio (Home)", [
    caso(
        "TC-HOME-001", "Inicio (/home)", "Vista general al cargar",
        "Sesión iniciada",
        [
            "Abre la pestaña Inicio.",
        ],
        "Muestra el saludo con el nombre del usuario, su avatar/iniciales, el código QR personal, y la lista "
        "de eventos en curso (o el estado vacío si no hay ninguno).",
    ),
    caso(
        "TC-HOME-002", "Inicio (/home)", "Nombre de usuario largo no rompe el diseño",
        "Cuenta de prueba con nombre y apellido largos (o usar una pantalla angosta)",
        [
            "Inicia sesión con un usuario cuyo nombre completo sea largo, o reduce el ancho de la ventana/usa "
            "un teléfono pequeño.",
            "Observa la barra superior de Inicio.",
        ],
        "El texto de saludo se ajusta de tamaño para caber en el espacio disponible, sin superponerse ni "
        "invadir el botón de notificaciones ni el de ajustes.",
    ),
    caso(
        "TC-HOME-003", "Inicio (/home)", "Refresco automático al iniciar un evento en curso",
        "Un evento programado con hora de inicio cercana (o forzada por un administrador), app abierta y en primer plano",
        [
            "Deja la app abierta en la pestaña Inicio.",
            "Espera a que un evento programado cambie automáticamente su estado a 'en curso' (o pide a un "
            "administrador que lo confirme desde el backend).",
            "No toques ni refresques la pantalla manualmente.",
        ],
        "La tarjeta del evento aparece sola en la sección 'eventos en curso' sin necesidad de recargar la app.",
    ),
    caso(
        "TC-HOME-004", "Inicio (/home)", "Reanudar la app desde segundo plano refresca los datos",
        "App con al menos un evento en curso, minimizada varios minutos",
        [
            "Minimiza la app (cambia a otra app o pestaña) por 2-3 minutos.",
            "Vuelve a la app.",
        ],
        "Al volver a primer plano, la lista de eventos en curso se actualiza sola (sin spinner de pantalla "
        "completa, solo actualización silenciosa).",
    ),
    caso(
        "TC-HOME-005", "Inicio (/home)", "Código QR personal es legible",
        "Sesión iniciada",
        [
            "Abre la tarjeta de tu código QR personal en Inicio.",
            "Con otro dispositivo (o la pantalla de escaneo de un administrador), escanea el QR mostrado.",
        ],
        "El QR se lee correctamente y corresponde a tu identidad de usuario.",
    ),
])

# 4. EVENTOS: LISTA, CREACIÓN, EDICIÓN, BORRADORES ----------------------------
modulo("4. Gestión de Eventos", [
    caso(
        "TC-EVT-001", "Eventos (/eventos)", "Listado de eventos",
        "Sesión iniciada",
        [
            "Abre la pestaña Eventos.",
        ],
        "Muestra la lista de eventos disponibles con su información básica (título, fecha, estado). Si no hay "
        "eventos, muestra un estado vacío claro en vez de una pantalla en blanco.",
    ),
    caso(
        "TC-EVT-002", "Crear evento — Paso 1 (/crear_evento)", "Completar información básica del evento",
        "Cuenta con permiso para crear eventos",
        [
            "Presiona 'Crear evento'.",
            "Completa título, descripción, lugar y tipo de evento.",
            "Presiona 'Siguiente'.",
        ],
        "Avanza al paso 2 solo si los campos obligatorios están completos; de lo contrario muestra los errores "
        "de validación correspondientes junto a cada campo.",
    ),
    caso(
        "TC-EVT-003", "Crear evento — Paso 2 (/crear_evento)", "Fecha, duración y audiencia",
        "Paso 1 completado",
        [
            "Selecciona fecha y hora de inicio y de fin.",
            "Configura la audiencia (general o dirigida por tags).",
            "Presiona 'Siguiente'.",
        ],
        "No permite continuar si la fecha/hora de fin es anterior a la de inicio. La sección de audiencia se "
        "muestra completa en este paso (no en el paso 3).",
    ),
    caso(
        "TC-EVT-004", "Crear evento — Paso 3 (/crear_evento)", "Revisión y publicación",
        "Pasos 1 y 2 completados",
        [
            "Revisa el resumen del evento.",
            "Presiona 'Guardar borrador'.",
            "Vuelve a entrar y ahora presiona 'Publicar evento'.",
        ],
        "'Guardar borrador' guarda sin notificar a nadie y aparece en la lista de Borradores. 'Publicar evento' "
        "cambia el estatus a programado y, si la audiencia es dirigida, dispara notificación a los usuarios "
        "cuyos tags coincidan.",
    ),
    caso(
        "TC-EVT-005", "Borradores (/borradores)", "Retomar y publicar un borrador",
        "Al menos un evento guardado como borrador",
        [
            "Abre la lista de Borradores.",
            "Selecciona un borrador existente.",
            "Complétalo y publícalo.",
        ],
        "El borrador se elimina de la lista de borradores y aparece como evento programado en la lista general.",
    ),
    caso(
        "TC-EVT-006", "Editar evento (/crear_evento/:id)", "Editar un evento existente",
        "Evento propio ya publicado, aún no iniciado",
        [
            "Desde el detalle de un evento, presiona 'Editar'.",
            "Modifica algún campo (ej. el lugar).",
            "Guarda los cambios.",
        ],
        "Los cambios se reflejan de inmediato en el listado y detalle del evento.",
    ),
    caso(
        "TC-EVT-007", "Crear evento", "Doble toque rápido en 'Publicar'",
        "Formulario de evento completo y válido",
        [
            "Completa el formulario de creación de evento.",
            "Toca el botón 'Publicar' dos veces muy rápido (simulando un doble-tap accidental).",
        ],
        "Se crea un único evento, no dos duplicados. El botón se deshabilita o ignora el segundo toque mientras "
        "se procesa el primero.",
    ),
])

# 5. PANEL DE CONTROL DE EVENTO -----------------------------------------------
modulo("5. Panel de Control de Evento", [
    caso(
        "TC-PANEL-001", "Panel de control (/eventos/:id/panel)", "Ver detalle y contador de asistencia en vivo",
        "Evento en curso con al menos un colaborador",
        [
            "Abre el panel de control de un evento en curso.",
            "Observa el contador de presentes.",
            "Registra la asistencia de alguien (desde otro dispositivo o el escáner) y vuelve a mirar el contador.",
        ],
        "El contador de presentes se actualiza en tiempo real sin necesidad de recargar la pantalla.",
    ),
    caso(
        "TC-PANEL-002", "Panel de control", "Cerrar evento manualmente",
        "Evento en curso, cuenta con permiso de panel de control",
        [
            "Presiona 'Cerrar evento'.",
            "Confirma la acción en el diálogo de confirmación.",
        ],
        "El evento cambia a estatus finalizado. Si tiene 'marcar ausentes automático' activado, los asistentes "
        "esperados que no registraron entrada quedan marcados como ausentes y reciben notificación.",
    ),
    caso(
        "TC-PANEL-003", "Buscar asistente (/eventos/:id/panel/buscar)", "Registro manual de asistencia por búsqueda",
        "Evento en curso con modo de registro manual habilitado",
        [
            "Abre 'Buscar asistente' desde el panel.",
            "Busca a una persona por nombre o cédula.",
            "Selecciónala y regístrala como presente.",
        ],
        "La persona aparece marcada como presente y el contador se actualiza. Buscar un texto sin resultados "
        "muestra un estado vacío claro, no un error.",
    ),
    caso(
        "TC-PANEL-004", "Colaboradores del evento (/eventos/:id/panel/colaboradores)", "Asignar y remover colaboradores",
        "Evento propio, al menos un usuario disponible para asignar",
        [
            "Abre 'Colaboradores' desde el panel.",
            "Asigna a un usuario como colaborador.",
            "Verifica que ese usuario recibe la notificación correspondiente.",
            "Remuévelo y verifica la notificación de remoción.",
        ],
        "El colaborador asignado gana acceso al panel de ese evento; al removerlo, pierde el acceso y ambos "
        "reciben su notificación correspondiente.",
    ),
    caso(
        "TC-PANEL-005", "Panel de control", "Exportar / copiar lista de asistentes",
        "Evento con al menos un asistente registrado",
        [
            "Desde el panel, usa la opción de copiar/exportar la lista de asistentes.",
        ],
        "Muestra confirmación de 'Lista copiada al portapapeles' (o equivalente) y el contenido copiado "
        "corresponde a los asistentes reales del evento.",
    ),
])

# 6. ESCANEO QR ----------------------------------------------------------------
modulo("6. Escaneo QR", [
    caso(
        "TC-QR-001", "Escanear (tab, /escanear)", "Autoregistro escaneando el QR del evento",
        "Evento en curso que permite QR de evento, cámara con permiso concedido",
        [
            "Toca la pestaña 'Escanear' en el menú inferior.",
            "Apunta la cámara al código QR del evento (mostrado en el panel de control o proyectado en el lugar).",
        ],
        "Detecta el QR, confirma tu asistencia y muestra un mensaje de éxito con tu nombre. Un segundo escaneo "
        "del mismo QR no debe duplicar el registro.",
    ),
    caso(
        "TC-QR-002", "Escanear usuario (/eventos/:id/panel/qr_usuario)", "Staff registra asistencia escaneando el QR de un asistente",
        "Cuenta con acceso al panel de control, evento que permite QR de usuario",
        [
            "Desde el panel de control, abre 'Escanear QR de usuario'.",
            "Escanea el QR personal de un asistente (pantalla de Inicio de esa persona).",
        ],
        "Muestra el nombre de la persona reconocida y confirma el registro de su entrada. Si ya estaba "
        "registrada, indica 'Ya registrado' en vez de duplicar la fila de asistencia.",
    ),
    caso(
        "TC-QR-003", "Escaneo QR (ambas pantallas)", "Sin permiso de cámara concedido",
        "Permiso de cámara del navegador/dispositivo denegado a propósito",
        [
            "Antes de abrir la app, deniega el permiso de cámara en la configuración del navegador/dispositivo.",
            "Abre cualquiera de las dos pantallas de escaneo.",
        ],
        "Muestra una pantalla de error/estado claro indicando que no hay acceso a la cámara, sin cerrar la app "
        "ni dejar una pantalla en blanco o congelada.",
    ),
    caso(
        "TC-QR-004", "Escaneo QR", "QR inválido o no reconocido",
        "Cámara con permiso concedido",
        [
            "Escanea un código QR que no corresponda a un usuario ni evento de la app (ej. un QR de otra web).",
        ],
        "Muestra un mensaje de 'código no válido' sin registrar nada ni bloquear la cámara para seguir intentando.",
    ),
    caso(
        "TC-QR-005", "Escaneo QR de usuario", "QR de carnet físico (cédula) vs. QR nativo de la app",
        "Un asistente con QR de carnet físico impreso (formato cédula, ej. 'V-12345678') y otro con el QR "
        "nativo de la app",
        [
            "Escanea el QR nativo de la app de un usuario (código UUID interno).",
            "Escanea el QR de un carnet físico universitario de otro usuario (con prefijo de cédula).",
        ],
        "Ambos formatos se reconocen correctamente y registran a la persona correcta — la app debe extraer el "
        "identificador limpio en ambos casos, sin importar prefijos o guiones del carnet.",
    ),
    caso(
        "TC-QR-006", "Escaneo QR", "Evento que no permite escaneo QR",
        "Evento configurado sin permitir QR de evento",
        [
            "Intenta autoregistrarte escaneando el QR de un evento que tiene deshabilitado el registro por QR.",
        ],
        "Muestra un mensaje claro de que ese evento no permite QR (ej. 'Este evento no permite...'), sin "
        "registrar asistencia.",
    ),
])

# 7. HISTORIAL -----------------------------------------------------------------
modulo("7. Historial (Mi historial)", [
    caso(
        "TC-HIST-001", "Historial (/historial)", "Ver historial de asistencia propio",
        "Usuario con al menos un evento asistido en el pasado",
        [
            "Abre la pestaña Historial.",
        ],
        "Muestra la lista de eventos pasados con tu estatus de asistencia (asistió, salió antes, ausente).",
    ),
    caso(
        "TC-HIST-002", "Historial", "Filtrar por texto y por rango de fechas",
        "Historial con varios eventos de distintas fechas",
        [
            "Escribe parte del título de un evento en el buscador.",
            "Limpia el buscador y en su lugar aplica un filtro de rango de fechas.",
        ],
        "La lista se filtra correctamente en ambos casos. Combinar texto + fecha filtra por ambos criterios a "
        "la vez. Una búsqueda sin resultados muestra un estado vacío, no un error.",
    ),
    caso(
        "TC-HIST-003", "Historial", "Pestañas de asistió / salió antes / ausente",
        "Historial con eventos en al menos dos de esas categorías",
        [
            "Cambia entre las pestañas de categoría dentro de Historial.",
        ],
        "Cada pestaña muestra solo los eventos correspondientes a esa categoría, con el conteo coherente.",
    ),
    caso(
        "TC-HIST-004", "Historial", "Exportar historial a PDF",
        "Historial con al menos un evento",
        [
            "Presiona el botón de exportar/compartir.",
            "Espera a que se genere el archivo.",
        ],
        "Genera y permite compartir/descargar un PDF legible con tu nombre y el detalle de tu historial. "
        "Muestra un indicador de carga mientras se genera, sin bloquear el resto de la pantalla indefinidamente.",
    ),
])

# 8. NOTIFICACIONES PUSH --------------------------------------------------------
modulo("8. Notificaciones Push", [
    caso(
        "TC-NOTIF-001", "Notificaciones (/notificaciones)", "Ver bandeja de notificaciones",
        "Usuario con al menos una notificación recibida",
        [
            "Toca el ícono de campana en la barra superior.",
        ],
        "Muestra la lista de notificaciones ordenadas de más reciente a más antigua, distinguiendo visualmente "
        "las leídas de las no leídas.",
    ),
    caso(
        "TC-NOTIF-002", "Notificaciones — badge", "Contador en tiempo real",
        "App abierta, capacidad de generar una notificación nueva (ej. que un admin te asigne un rol)",
        [
            "Deja la app abierta en cualquier pantalla, sin entrar a Notificaciones.",
            "Haz que llegue una notificación nueva (pide a otra persona que dispare una acción que te notifique).",
        ],
        "El número rojo sobre el ícono de campana aparece o se incrementa solo, sin recargar la app.",
    ),
    caso(
        "TC-NOTIF-003", "Notificaciones", "Marcar como leída (individual y todas)",
        "Al menos dos notificaciones no leídas",
        [
            "Toca una notificación no leída individual.",
            "Presiona 'Marcar todas' para el resto.",
        ],
        "La notificación tocada cambia su estilo a leída de inmediato. 'Marcar todas' limpia el contador del "
        "badge a cero.",
    ),
    caso(
        "TC-NOTIF-004", "Push — app cerrada/en segundo plano", "Recepción de notificación con la app cerrada",
        "Notificaciones activadas en el dispositivo, app cerrada o minimizada",
        [
            "Cierra completamente la app (o pásala a segundo plano).",
            "Haz que llegue una notificación (ej. que se cree un evento que te aplica).",
        ],
        "El sistema operativo muestra la notificación push nativa (banner/sonido según configuración del "
        "dispositivo), incluso con la app cerrada.",
    ),
    caso(
        "TC-NOTIF-005", "Push — clic en notificación", "Abrir la app desde la notificación del sistema",
        "Notificación push recibida y visible en el centro de notificaciones del sistema",
        [
            "Toca la notificación push recibida en la bandeja del sistema operativo.",
        ],
        "Abre (o trae a primer plano) la app. Idealmente navega directo al contenido relacionado (ej. el evento "
        "mencionado) en vez de solo abrir en Inicio.",
    ),
    caso(
        "TC-NOTIF-006", "Push — app en primer plano", "Aviso visual al recibir push con la app abierta",
        "App abierta y en uso activo, notificaciones activadas",
        [
            "Mantén la app abierta y visible.",
            "Haz que llegue una notificación push.",
        ],
        "Aparece un aviso flotante dentro de la app (no un banner del sistema, ya que la app está en primer "
        "plano) mostrando el contenido de la notificación, y el badge de la campana se actualiza a la vez.",
    ),
    caso(
        "TC-NOTIF-007", "Notificaciones — banner de activación", "Activar notificaciones desde el banner",
        "Cuenta sin token de notificaciones push registrado (dispositivo/navegador nuevo)",
        [
            "Entra a Notificaciones con una cuenta/dispositivo que nunca activó los push.",
            "Presiona 'Activar' en el banner que aparece.",
            "Acepta el permiso del navegador/sistema cuando se solicite.",
        ],
        "El banner desaparece tras activarse correctamente, y a partir de ese momento el dispositivo puede "
        "recibir push reales (verificable repitiendo TC-NOTIF-004).",
    ),
])

# 9. FORMULARIOS Y VALIDACIONES (TRANSVERSAL) ----------------------------------
modulo("9. Formularios y Validaciones (transversal)", [
    caso(
        "TC-FORM-001", "Cualquier formulario con campos de texto", "Campos vacíos obligatorios",
        "Cualquier pantalla con un formulario (crear evento, crear usuario, crear rol/tag/tipo de evento, etc.)",
        [
            "Abre un formulario de creación.",
            "Deja vacíos los campos obligatorios.",
            "Intenta enviar/guardar.",
        ],
        "Bloquea el envío y resalta con un mensaje claro cuáles campos faltan, sin llamar al servidor "
        "innecesariamente.",
    ),
    caso(
        "TC-FORM-002", "Cualquier formulario con campos de texto", "Texto extremadamente largo",
        "Cualquier formulario con campos de texto libre (título de evento, descripción, nombre, etc.)",
        [
            "Pega un texto muy largo (varios párrafos) en un campo pensado para un texto corto (ej. título).",
            "Intenta guardar.",
        ],
        "El campo limita la longitud visualmente (o el formulario rechaza con un mensaje claro) sin romper el "
        "diseño de la pantalla ni provocar un error críptico del servidor.",
    ),
    caso(
        "TC-FORM-003", "Cualquier formulario con campos de texto", "Caracteres especiales y emojis",
        "Cualquier formulario de texto libre",
        [
            "Ingresa caracteres especiales (comillas, símbolos, acentos) y emojis en campos de texto.",
            "Guarda y vuelve a abrir el registro creado.",
        ],
        "El contenido se guarda y se muestra exactamente igual a como se ingresó, sin errores de codificación "
        "ni caracteres corruptos.",
    ),
    caso(
        "TC-FORM-004", "Cualquier botón de acción (crear, guardar, eliminar, publicar)", "Doble toque rápido",
        "Formulario completo y válido, listo para enviar",
        [
            "Completa cualquier formulario de creación/edición.",
            "Toca el botón de acción principal dos veces muy rápido, simulando un doble toque accidental.",
        ],
        "Se ejecuta la acción una sola vez (no se crean duplicados ni se envían dos peticiones). El botón se "
        "deshabilita visualmente mientras procesa.",
    ),
    caso(
        "TC-FORM-005", "Formularios con selección de fecha/hora", "Fechas y horas ilógicas",
        "Cualquier formulario con selectores de fecha/hora de inicio y fin",
        [
            "Selecciona una fecha/hora de fin anterior a la de inicio.",
            "Intenta guardar.",
        ],
        "Bloquea el guardado y explica el problema claramente, sin permitir un registro con fechas inconsistentes.",
    ),
    caso(
        "TC-FORM-006", "Formularios con confirmación (contraseña, eliminar)", "Confirmación de campos sensibles",
        "Formulario de registro o cambio de contraseña",
        [
            "Ingresa una contraseña y su confirmación con valores distintos.",
            "Intenta continuar.",
        ],
        "Bloquea el envío indicando que las contraseñas no coinciden.",
    ),
])

# 10. RENDIMIENTO VISUAL Y UX (TRANSVERSAL) -------------------------------------
modulo("10. Rendimiento Visual y UX (transversal)", [
    caso(
        "TC-UX-001", "Cualquier pantalla con carga de datos", "Indicador de carga visible",
        "Conexión de red normal",
        [
            "Entra a una pantalla que carga datos del servidor (ej. Eventos, Historial, Estadísticas).",
            "Observa el instante inicial antes de que aparezcan los datos.",
        ],
        "Muestra un spinner/esqueleto de carga claro mientras espera la respuesta — nunca una pantalla en "
        "blanco sin ninguna indicación.",
    ),
    caso(
        "TC-UX-002", "Cualquier pantalla con carga de datos", "Comportamiento con red muy lenta",
        "Throttling de red activado (ej. 'Slow 3G' en DevTools)",
        [
            "Activa un throttling de red lento.",
            "Navega a una pantalla que dependa de datos del servidor.",
        ],
        "El indicador de carga se mantiene visible todo el tiempo que dure la espera, sin que la interfaz "
        "parezca congelada o sin respuesta. Si tarda demasiado, debería existir algún mensaje o tiempo de espera.",
    ),
    caso(
        "TC-UX-003", "Cualquier pantalla con carga de datos", "Comportamiento sin conexión a internet",
        "Dispositivo en modo avión o sin red",
        [
            "Activa el modo avión.",
            "Intenta navegar o realizar una acción que requiera el servidor (ej. crear un evento).",
        ],
        "Muestra un mensaje claro de falta de conexión en vez de un error técnico críptico o una pantalla en "
        "blanco. Al recuperar la conexión, la app se recupera sin necesidad de recargarla por completo.",
    ),
    caso(
        "TC-UX-004", "Listas largas (Eventos, Historial, Usuarios)", "Fluidez del scroll",
        "Un listado con suficientes elementos para requerir scroll",
        [
            "Desplázate rápido hacia abajo y hacia arriba en una lista larga.",
        ],
        "El scroll se siente fluido, sin saltos ni tirones perceptibles, incluso mientras cargan más elementos "
        "(scroll infinito, si aplica).",
    ),
    caso(
        "TC-UX-005", "Navegación entre pantallas", "Transiciones y animaciones",
        "Cualquier navegación entre pantallas (push/pop, cambio de pestaña)",
        [
            "Navega hacia una pantalla interna y regresa.",
            "Cambia entre pestañas del menú inferior varias veces seguidas.",
        ],
        "Las transiciones se ven suaves y consistentes con el resto de la app, sin parpadeos, contenido que "
        "'salta' de posición, ni animaciones que se cortan a la mitad.",
    ),
    caso(
        "TC-UX-006", "Cualquier pantalla", "Rotación de pantalla",
        "Dispositivo con rotación automática habilitada",
        [
            "En una pantalla con formulario largo (ej. crear evento) y en el escáner QR, rota el dispositivo "
            "de vertical a horizontal y viceversa.",
        ],
        "El contenido se reacomoda sin cortarse, sin elementos superpuestos, y sin perder los datos ya "
        "ingresados en el formulario.",
    ),
    caso(
        "TC-UX-007", "Toda la app", "Consistencia visual entre pantallas",
        "Recorrido general por varias pantallas",
        [
            "Navega por al menos 8-10 pantallas distintas de la app.",
            "Compara colores, tipografía, espaciados y estilo de botones entre ellas.",
        ],
        "El diseño se siente consistente (misma paleta de colores, mismos componentes reutilizados) en toda "
        "la app, sin pantallas que 'desentonen' visualmente del resto.",
    ),
    caso(
        "TC-UX-008", "PWA — actualización de versión", "Detección de nueva versión con la app abierta",
        "Es posible desplegar una nueva versión de la app durante la prueba (coordinar con el equipo de "
        "desarrollo), pestaña/PWA abierta desde antes del despliegue",
        [
            "Deja la app abierta desde antes de que se publique una nueva versión.",
            "Espera a que el equipo de desarrollo despliegue una actualización.",
            "Sin cerrar la pestaña/app, intenta seguir usándola con normalidad.",
        ],
        "La app detecta la nueva versión y ofrece recargar (o lo hace sola de forma no disruptiva), en vez de "
        "quedarse atascada indefinidamente en una versión vieja con el riesgo de que deje de ser compatible "
        "con el backend.",
    ),
    caso(
        "TC-UX-009", "PWA — instalación", "Instalar la app como PWA en el dispositivo",
        "Navegador compatible con instalación de PWA (Chrome/Edge en Android o escritorio; Safari en iOS "
        "usando 'Compartir → Añadir a inicio')",
        [
            "Desde el navegador, sigue el flujo de instalación de la app a la pantalla de inicio/escritorio.",
            "Abre la app ya instalada desde su ícono.",
            "En iOS específicamente, verifica si las notificaciones push funcionan tras instalarla (requiere "
            "iOS 16.4+; en versiones anteriores no deberían prometerse como disponibles).",
        ],
        "La instalación se completa sin errores, el ícono y nombre de la app se ven correctos, y la app abre "
        "en modo standalone (sin la barra de direcciones del navegador). El comportamiento de push en iOS "
        "coincide con las limitaciones conocidas de la plataforma, sin fallar de forma confusa.",
    ),
    caso(
        "TC-UX-010", "Eventos — horarios", "Zona horaria del dispositivo distinta a la del servidor",
        "Evento de prueba con hora de inicio conocida; dispositivo con zona horaria distinta a "
        "'America/Caracas' configurada manualmente",
        [
            "Cambia la zona horaria del dispositivo a una distinta de Venezuela (ej. España o México).",
            "Abre un evento con hora de inicio conocida y compárala con la hora real de Caracas para ese evento.",
        ],
        "La hora mostrada en la app sigue correspondiendo al horario real del evento (convertida correctamente "
        "a la zona horaria local del dispositivo), no se desfasa ni muestra la hora de Caracas como si fuera "
        "la hora local sin convertir.",
    ),
    caso(
        "TC-UX-011", "Toda la app", "Accesibilidad básica",
        "Dispositivo con lector de pantalla disponible (TalkBack en Android, VoiceOver en iOS) o navegador "
        "con zoom/tamaño de fuente del sistema aumentado",
        [
            "Activa el lector de pantalla del sistema y navega por 3-4 pantallas clave (Login, Inicio, "
            "crear evento).",
            "Desactívalo y en su lugar aumenta el tamaño de fuente del sistema al máximo.",
        ],
        "El lector de pantalla anuncia los elementos interactivos de forma comprensible (botones, campos). "
        "Con la fuente del sistema aumentada, el texto crece sin que los botones queden inalcanzables ni el "
        "contenido se corte de forma ilegible.",
    ),
    caso(
        "TC-UX-012", "Panel de control de evento", "Edición concurrente del mismo evento",
        "Dos cuentas con permiso sobre el mismo evento, dos dispositivos",
        [
            "Con el Dispositivo A, abre la edición de un evento y cambia un campo (sin guardar aún).",
            "Con el Dispositivo B, edita ese mismo evento y guarda un cambio distinto.",
            "En el Dispositivo A, guarda tu cambio.",
        ],
        "El resultado final es predecible (normalmente el último guardado gana) y ninguno de los dos "
        "dispositivos muestra un error críptico o deja el evento en un estado inconsistente/corrupto.",
    ),
    caso(
        "TC-UX-013", "Escaneo QR", "Doble escaneo casi simultáneo del mismo asistente",
        "Dos operadores con acceso al panel de un mismo evento, cada uno con su propio escáner",
        [
            "Con dos dispositivos distintos, escanea el QR del mismo asistente casi al mismo tiempo (diferencia "
            "de menos de un segundo).",
        ],
        "Solo se registra una entrada de asistencia para esa persona — el segundo escaneo debe reconocer que "
        "ya fue registrada, sin crear una fila duplicada ni un conteo de presentes inflado.",
    ),
    caso(
        "TC-UX-014", "Listas con muchos datos (Eventos, Usuarios, Historial)", "Rendimiento con volumen alto de datos",
        "Un entorno de prueba con varios cientos de registros en alguna lista (eventos, usuarios o historial) "
        "— coordinar con el equipo de desarrollo para tener datos de prueba a esa escala",
        [
            "Abre una lista con varios cientos de elementos.",
            "Haz scroll hasta el final y busca/filtra dentro de esa lista.",
        ],
        "La lista carga de forma progresiva (paginación o scroll infinito) sin intentar traer todo de una vez "
        "de forma perceptible como congelamiento, y la búsqueda/filtro sigue respondiendo con fluidez.",
    ),
])

# 11. PERFIL --------------------------------------------------------------------
modulo("11. Perfil", [
    caso(
        "TC-PERFIL-001", "Perfil (/perfil)", "Ver y editar información propia",
        "Sesión iniciada",
        [
            "Abre la pestaña Perfil.",
            "Edita algún dato editable (ej. teléfono).",
            "Guarda los cambios.",
        ],
        "Los cambios se guardan y se reflejan de inmediato en la pantalla y en cualquier otro lugar donde se "
        "muestre ese dato (ej. Inicio).",
    ),
    caso(
        "TC-PERFIL-002", "Perfil", "Cerrar sesión",
        "Sesión iniciada",
        [
            "Presiona 'Cerrar sesión'.",
            "Confirma si se solicita confirmación.",
        ],
        "Vuelve a la pantalla de Login y no permite regresar a pantallas internas con el botón atrás del "
        "navegador/sistema.",
    ),
])

# 12. ADMINISTRACIÓN — USUARIOS --------------------------------------------------
modulo("12. Administración — Usuarios", [
    caso(
        "TC-USR-001", "Gestión de usuarios (/gestion_usuarios)", "Listado y búsqueda de usuarios",
        "Cuenta con permiso de administración de usuarios",
        [
            "Abre 'Gestión de usuarios'.",
            "Busca a un usuario por nombre.",
        ],
        "Muestra el listado completo y filtra correctamente al buscar. Una búsqueda sin resultados muestra "
        "un estado vacío claro.",
    ),
    caso(
        "TC-USR-002", "Editar usuario (/gestion_usuarios/:id/editar)", "Editar datos de un usuario",
        "Permiso de administración de usuarios",
        [
            "Abre un usuario desde el listado.",
            "Edita alguno de sus datos.",
            "Guarda los cambios.",
        ],
        "Los cambios quedan reflejados en el listado y en el perfil de ese usuario.",
    ),
    caso(
        "TC-USR-003", "Gestionar tags de usuario (/gestion_usuarios/:id/tags)", "Asignar/quitar tags a un usuario",
        "Permiso de administración de usuarios, al menos un tag creado",
        [
            "Abre la gestión de tags de un usuario.",
            "Asigna un tag y guarda.",
            "Quita ese mismo tag y guarda.",
        ],
        "Los tags asignados afectan qué eventos dirigidos ve/recibe ese usuario (verificable cruzando con "
        "TC-EVT-004 usando audiencia dirigida por ese tag).",
    ),
    caso(
        "TC-USR-004", "Gestionar roles de usuario (/gestion_usuarios/:id/roles)", "Asignar/quitar roles del sistema",
        "Permiso de administración de usuarios, al menos un rol creado",
        [
            "Abre la gestión de roles de un usuario.",
            "Asígnale un rol y guarda.",
            "Verifica que el usuario recibe la notificación de rol asignado.",
        ],
        "El usuario gana los permisos asociados a ese rol de inmediato (verificable pidiéndole que intente "
        "acceder a una sección antes bloqueada) y recibe su notificación.",
    ),
    caso(
        "TC-USR-005", "Revisión de usuarios (/revision_usuarios)", "Aprobar, rechazar y suspender cuentas",
        "Permiso de revisión de usuarios, al menos una cuenta pendiente",
        [
            "Abre 'Revisión de usuarios'.",
            "Aprueba una cuenta pendiente.",
            "Rechaza otra cuenta pendiente.",
            "Suspende una cuenta ya activa.",
        ],
        "Cada acción cambia el estatus correctamente y dispara la notificación correspondiente (cuenta "
        "aprobada / solicitud rechazada / cuenta suspendida) al usuario afectado.",
    ),
    caso(
        "TC-USR-006", "Ver perfil de usuario (/gestion_usuarios/:id/perfil)", "Ver el perfil completo de otro usuario",
        "Permiso de administración de usuarios",
        [
            "Desde el listado de 'Gestión de usuarios', abre un usuario y entra a su vista de perfil "
            "(distinta de 'Editar').",
        ],
        "Muestra los datos completos del usuario (contacto, tags, roles, estatus) en modo solo lectura, sin "
        "mezclarse con la pantalla de edición ni permitir cambios accidentales.",
    ),
])

# 13. ADMINISTRACIÓN — ROLES Y PERMISOS ------------------------------------------
modulo("13. Administración — Roles y Permisos", [
    caso(
        "TC-ROL-001", "Gestión de roles (/gestion_roles)", "Listado de roles",
        "Permiso de administración de roles",
        [
            "Abre 'Gestión de roles'.",
        ],
        "Muestra el listado de roles existentes con sus permisos asociados visibles o accesibles.",
    ),
    caso(
        "TC-ROL-002", "Crear/editar rol (/crear_rol, /crear_rol/:id)", "Crear un rol nuevo con permisos",
        "Permiso de administración de roles",
        [
            "Presiona 'Crear rol'.",
            "Asigna un nombre y selecciona uno o más permisos.",
            "Guarda.",
        ],
        "El rol se crea y queda disponible para asignar a usuarios (TC-USR-004) con exactamente los permisos "
        "seleccionados.",
    ),
    caso(
        "TC-ROL-003", "Permisos del sistema (/permisos)", "Ver catálogo de permisos",
        "Permiso de administración de permisos",
        [
            "Abre 'Permisos del sistema'.",
        ],
        "Muestra el catálogo completo de permisos disponibles en la app, legible y organizado.",
    ),
    caso(
        "TC-ROL-004", "Gestión de roles", "Estado vacío sin roles creados",
        "Proyecto nuevo o de prueba sin roles creados aún",
        [
            "Abre 'Gestión de roles' cuando todavía no existe ningún rol en el sistema.",
        ],
        "Muestra un estado vacío claro invitando a crear el primer rol, no una tabla en blanco ni un error.",
    ),
])

# 14. ADMINISTRACIÓN — TAGS -------------------------------------------------------
modulo("14. Administración — Tags", [
    caso(
        "TC-TAG-001", "Gestión de tags (/gestion_tags)", "Listado de tags",
        "Permiso de administración de tags",
        [
            "Abre 'Gestión de tags'.",
        ],
        "Muestra el listado de tags existentes, agrupados o identificados por su grupo/categoría si aplica.",
    ),
    caso(
        "TC-TAG-002", "Crear/editar tag (/crear_tag, /crear_tag/:id)", "Crear un tag nuevo",
        "Permiso de administración de tags",
        [
            "Presiona 'Crear tag'.",
            "Completa nombre y grupo del tag.",
            "Guarda.",
        ],
        "El tag queda disponible para asignar a usuarios (TC-USR-003) y para configurar audiencia dirigida "
        "de eventos (TC-EVT-003).",
    ),
    caso(
        "TC-TAG-003", "Gestión de tags", "Estado vacío sin tags creados",
        "Proyecto nuevo o de prueba sin tags creados aún",
        [
            "Abre 'Gestión de tags' cuando todavía no existe ningún tag en el sistema.",
        ],
        "Muestra un estado vacío claro invitando a crear el primer tag, no una tabla en blanco ni un error.",
    ),
])

# 15. ADMINISTRACIÓN — TIPOS DE EVENTO --------------------------------------------
modulo("15. Administración — Tipos de Evento", [
    caso(
        "TC-TIPO-001", "Gestión de tipos de evento (/gestion_tipos_evento)", "Listado de tipos de evento",
        "Permiso de administración de tipos de evento",
        [
            "Abre 'Gestión de tipos de evento'.",
        ],
        "Muestra el listado de tipos existentes (ej. Conferencia, Taller, etc.).",
    ),
    caso(
        "TC-TIPO-002", "Crear/editar tipo de evento", "Crear un tipo de evento nuevo",
        "Permiso de administración de tipos de evento",
        [
            "Presiona 'Crear tipo de evento'.",
            "Completa el nombre.",
            "Guarda.",
        ],
        "El tipo queda disponible en el selector del Paso 1 de creación de eventos (TC-EVT-002).",
    ),
    caso(
        "TC-TIPO-003", "Gestión de tipos de evento", "Estado vacío sin tipos creados",
        "Proyecto nuevo o de prueba sin tipos de evento creados aún",
        [
            "Abre 'Gestión de tipos de evento' cuando todavía no existe ningún tipo en el sistema.",
        ],
        "Muestra un estado vacío claro invitando a crear el primer tipo, no una tabla en blanco ni un error. "
        "El Paso 1 de crear evento (TC-EVT-002) debe manejar con gracia el caso de no tener ningún tipo "
        "disponible para seleccionar.",
    ),
])

# 16. ESTADÍSTICAS -----------------------------------------------------------------
modulo("16. Estadísticas", [
    caso(
        "TC-STAT-001", "Estadísticas (/estadisticas)", "Ver dashboard general",
        "Permiso de estadísticas, datos históricos disponibles",
        [
            "Abre 'Estadísticas'.",
        ],
        "Carga las gráficas y resúmenes numéricos sin errores, con valores consistentes con los datos reales "
        "conocidos (ej. número total de eventos).",
    ),
    caso(
        "TC-STAT-002", "Estadísticas", "Aplicar filtros (fecha, tipo, creador, tag)",
        "Permiso de estadísticas",
        [
            "Aplica un filtro de rango de fechas.",
            "Combina con un filtro de tipo de evento.",
        ],
        "Todas las gráficas y números se recalculan según los filtros aplicados, de forma coherente entre sí.",
    ),
    caso(
        "TC-STAT-003", "Estadísticas", "Sin datos históricos (o filtro sin resultados)",
        "Cuenta/proyecto sin eventos pasados, o un filtro aplicado que no coincide con ningún dato",
        [
            "Abre 'Estadísticas' en un entorno sin eventos históricos, o aplica un filtro de fecha muy "
            "restrictivo que no tenga datos.",
        ],
        "Muestra un estado vacío claro en cada gráfica/sección (ej. 'Sin datos para este período'), sin "
        "gráficas rotas, valores 'NaN' o divisiones por cero visibles.",
    ),
])

# 17. AUDITORÍA DE EVENTO -----------------------------------------------------------
modulo("17. Auditoría de Evento", [
    caso(
        "TC-AUD-001", "Auditoría de evento (/auditoria_evento)", "Ver historial/log de cambios de un evento",
        "Permiso de auditoría, evento con historial de cambios",
        [
            "Abre la auditoría de un evento con varios cambios registrados (creación, ediciones, cierre).",
        ],
        "Muestra el registro cronológico de cambios con quién y cuándo los hizo, de forma legible.",
    ),
])

# 18. REGISTRO DE VISITANTES EXTERNOS (FORÁNEOS) --------------------------------
# Funcionalidad completa encontrada en el código (modal_foraneo.dart,
# permite_foraneos en eventos, registrarForaneo en EventosEnCursoCubit) que
# no tenía cobertura en la primera versión de este plan.
modulo("18. Registro de Visitantes Externos (Foráneos)", [
    caso(
        "TC-FORANEO-001", "Tarjeta de evento en curso (Inicio)", "Registrar un visitante externo sin cuenta",
        "Evento en curso con la opción 'permite foráneos' habilitada, cuenta con permiso para registrar "
        "asistencia en ese evento",
        [
            "En Inicio, sobre la tarjeta del evento en curso, presiona el botón 'Usuario foráneo'.",
            "Completa primer nombre, primer apellido y cédula (contacto es opcional).",
            "Presiona registrar/guardar.",
        ],
        "Registra la asistencia del visitante externo sin necesidad de que tenga una cuenta en el sistema, y "
        "el contador de presentes del evento se actualiza para incluirlo.",
    ),
    caso(
        "TC-FORANEO-002", "Tarjeta de evento en curso (Inicio)", "Evento que NO permite foráneos",
        "Evento en curso con la opción 'permite foráneos' deshabilitada",
        [
            "Abre la tarjeta de un evento en curso que no permite foráneos.",
            "Busca el botón 'Usuario foráneo'.",
        ],
        "El botón no aparece (o aparece deshabilitado con una explicación) — no debe ser posible abrir el "
        "modal de registro de foráneo para este evento.",
    ),
    caso(
        "TC-FORANEO-003", "Modal de registro de foráneo", "Campos obligatorios vacíos",
        "Modal de 'Usuario foráneo' abierto",
        [
            "Abre el modal 'Usuario foráneo'.",
            "Deja vacío el campo de cédula (o nombre/apellido).",
            "Intenta registrar.",
        ],
        "No permite guardar mientras falten nombre, apellido o cédula — el campo de contacto es el único "
        "opcional. El modal no se cierra ni muestra un error críptico, simplemente no procede.",
    ),
])


# ─── Construcción del documento ──────────────────────────────────────────────

def main():
    doc = Document()

    estilo_normal = doc.styles["Normal"]
    estilo_normal.font.name = "Calibri"
    estilo_normal.font.size = Pt(10.5)
    estilo_normal.font.color.rgb = COLOR_TEXTO

    for seccion in doc.sections:
        seccion.left_margin = Cm(2)
        seccion.right_margin = Cm(2)
        seccion.top_margin = Cm(1.8)
        seccion.bottom_margin = Cm(1.8)

    agregar_encabezado_portada(doc)
    agregar_instrucciones(doc)
    agregar_indice_modulos(doc, MODULOS)

    for nombre, casos in MODULOS:
        agregar_modulo(doc, nombre, casos)

    agregar_resumen_hallazgos(doc)

    salida = "Plan_de_Pruebas_Manuales_QA.docx"
    doc.save(salida)

    total_casos = sum(len(c) for _, c in MODULOS)
    print(f"OK: '{salida}' generado con {len(MODULOS)} módulos y {total_casos} casos de prueba.")


if __name__ == "__main__":
    main()
